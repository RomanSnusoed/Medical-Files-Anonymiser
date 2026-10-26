"""Flask endpoint tests without downloading large NLP models."""
from __future__ import annotations

import io
from pathlib import Path

import pytest
from docx import Document

import app as web


@pytest.fixture()
def client(monkeypatch):
    web.app.config.update(TESTING=True, SECRET_KEY="test-key", MAX_CONTENT_LENGTH=web.MAX_UPLOAD_BYTES)
    web.app.debug = False
    monkeypatch.delenv("LOG_PII", raising=False)
    monkeypatch.delenv("ENABLE_EXTERNAL_AI", raising=False)
    with web._jobs_lock:
        web._jobs.clear()
    with web.app.test_client() as test_client:
        test_client.get("/")
        yield test_client
    with web._jobs_lock:
        web._jobs.clear()


def _csrf(client):
    with client.session_transaction() as state:
        return state["csrf"]


def _docx_upload():
    doc = Document()
    doc.add_paragraph("Synthetic patient Alice visited Berlin.")
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return (output, "patient.docx")


def test_csrf_is_required(client):
    response = client.post("/process", data={"file": _docx_upload()})
    assert response.status_code == 400


def test_wrong_file_format_is_rejected(client):
    response = client.post(
        "/process",
        data={
            "csrf_token": _csrf(client),
            "file": (io.BytesIO(b"not-a-document"), "fake.docx"),
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert b"Invalid DOCX" in response.data


def test_process_download_is_session_bound(client, monkeypatch):
    monkeypatch.setattr(web, "detect", lambda text: "en")
    monkeypatch.setattr(
        web, "_anonymize",
        lambda text, source, target, verbose: (
            "Synthetic patient [REDACTED] visited Berlin.",
            "Synthetic patient [REDACTED] visited Berlin.",
            "en", ["Presidio: PERSON at [18:23]"], ["Detector pass completed."],
        ),
    )
    monkeypatch.setattr(web, "build_pdf_bytes", lambda text: b"%PDF-test-only")
    response = client.post(
        "/process",
        data={"csrf_token": _csrf(client), "file": _docx_upload(), "generate_pdf": "on"},
        content_type="multipart/form-data",
    )
    assert response.status_code == 303
    page = client.get(response.headers["Location"])
    assert page.status_code == 200
    assert b"[REDACTED]" in page.data
    assert b"patient Alice" not in page.data
    token = response.headers["Location"].split("/")[-1]
    downloaded = client.get(f"/download/{token}")
    assert downloaded.status_code == 200
    assert downloaded.data == b"%PDF-test-only"
    assert "attachment" in downloaded.headers["Content-Disposition"]
    assert downloaded.headers["Cache-Control"].startswith("no-store")

    with web.app.test_client() as stranger:
        assert stranger.get(f"/download/{token}").status_code == 404


def test_pdf_checkbox_respected(client, monkeypatch):
    monkeypatch.setattr(web, "detect", lambda text: "en")
    monkeypatch.setattr(
        web, "_anonymize",
        lambda text, source, target, verbose: (text, text, "en", [], []),
    )
    response = client.post(
        "/process",
        data={"csrf_token": _csrf(client), "file": _docx_upload()},
        content_type="multipart/form-data",
    )
    assert response.status_code == 303
    token = response.headers["Location"].split("/")[-1]
    assert client.get(f"/download/{token}").status_code == 404


def test_api_requires_both_env_and_consent(client, monkeypatch):
    monkeypatch.setattr(web, "detect", lambda text: "en")
    monkeypatch.setattr(
        web, "_anonymize",
        lambda text, source, target, verbose: ("[REDACTED]", "[REDACTED]", "en", [], []),
    )
    response = client.post(
        "/process",
        data={"csrf_token": _csrf(client), "file": _docx_upload()},
        content_type="multipart/form-data",
    )
    token = response.headers["Location"].split("/")[-1]

    def should_not_run(_text):
        pytest.fail("API must remain disabled")

    monkeypatch.setattr(web, "generate_explanation", should_not_run)
    response = client.post(
        "/generate_explanation",
        data={
            "csrf_token": _csrf(client),
            "job_token": token,
            "external_api_consent": "yes",
        },
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert b"disabled by default" in response.data


def test_docx_extraction_includes_tables_and_headers(tmp_path):
    from extractors import extract_text_from_docx

    file = Path(tmp_path) / "example.docx"
    doc = Document()
    doc.add_paragraph("Main text")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Patient ID: ZZ-2020"
    doc.sections[0].header.paragraphs[0].text = "Hospital Name"
    doc.save(file)
    text = extract_text_from_docx(file)
    assert "Main text" in text
    assert "Patient ID: ZZ-2020" in text
    assert "Hospital Name" in text


def test_pdf_extracts_text_and_scanned_page(tmp_path, monkeypatch):
    import fitz
    import extractors

    file = Path(tmp_path) / "mix.pdf"
    with fitz.open() as document:
        page = document.new_page()
        page.insert_text((75, 75), "Typed PDF text")
        document.new_page()
        document.save(file)

    monkeypatch.setattr(extractors, "ensure_tesseract_cmd", lambda: "tesseract")
    monkeypatch.setattr(extractors.pytesseract, "image_to_string", lambda *_args, **_kwargs: "OCR test text")
    output = extractors.extract_text_from_pdf(file)
    assert "Typed PDF text" in output
    assert "OCR test text" in output
