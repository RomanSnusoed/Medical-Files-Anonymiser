"""File extraction. All processing occurs in memory where feasible."""
from __future__ import annotations

from pathlib import Path

import fitz
import pytesseract
from docx import Document
from PIL import Image, ImageOps

from utils import ensure_tesseract_cmd, preprocess_image


def extract_text_from_image(image_path: str | Path) -> str:
    ensure_tesseract_cmd()
    image = preprocess_image(image_path)
    return pytesseract.image_to_string(image, config="--oem 3 --psm 6").strip()


def extract_text_from_pdf(pdf_path: str | Path) -> str:
    """OCR each page with no readable text; do not persist embedded images."""
    texts = []
    with fitz.open(pdf_path) as document:
        if document.needs_pass:
            raise ValueError("Password-protected PDFs are not supported")
        if len(document) > 50:
            raise ValueError("PDF exceeds 50-page limit")
        for page in document:
            text = page.get_text("text")
            if not text.strip():
                ensure_tesseract_cmd()
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                text = pytesseract.image_to_string(ImageOps.grayscale(image), config="--psm 3")
            texts.append(text)
    return "\n".join(texts)


def extract_text_from_docx(docx_path: str | Path) -> str:
    """Include paragraphs, table cells, headers and footers."""
    doc = Document(docx_path)
    paragraphs = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))
    for section in doc.sections:
        paragraphs.extend(p.text for p in section.header.paragraphs)
        paragraphs.extend(p.text for p in section.footer.paragraphs)
    return "\n".join(part for part in paragraphs if part.strip())
